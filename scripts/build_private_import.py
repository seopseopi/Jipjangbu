#!/usr/bin/env python3
"""Build a one-time private D1 migration from an Excel-exported .xlsx file."""

from __future__ import annotations

import argparse
import hashlib
from collections import OrderedDict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - operator-facing setup error
    raise SystemExit("openpyxl이 필요합니다: python3 -m pip install openpyxl") from exc


EVENT_WORK_TYPES = {
    "매물등록", "매물수정", "타계약확인", "매물취소", "가계약", "계약서작성",
    "중도금", "잔금", "계약파기", "계약취소", "경매확인",
}
REQUIRED_SHEETS = {"업무일지", "업무일지_과거", "고객목록", "매물목록"}
BREAK = "\n--> statement-breakpoint\n"


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).replace("_x000D_", "\n").replace("\r\n", "\n").replace("\r", "\n").strip()


def date_text(value: Any) -> str | None:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    text = clean(value)
    if not text or text == "NO매물":
        return None
    return text[:10].replace("/", "-")


def timestamp_text(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d 00:00:00")
    text = clean(value)
    return text or None


def listing_key(values: Iterable[Any]) -> str:
    return "|".join(clean(value).lower() for value in values)


def literal(value: Any) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, (int, float)):
        return str(value)
    return "'" + str(value).replace("'", "''") + "'"


def insert_chunks(table: str, columns: list[str], rows: list[tuple[Any, ...]], chunk_size: int = 40) -> list[str]:
    statements: list[str] = []
    quoted_columns = ", ".join(f'"{column}"' for column in columns)
    for start in range(0, len(rows), chunk_size):
        values = ",\n".join(
            "(" + ", ".join(literal(value) for value in row) + ")"
            for row in rows[start:start + chunk_size]
        )
        statements.append(f'INSERT INTO "{table}" ({quoted_columns}) VALUES\n{values};')
    return statements


def build_customers(workbook) -> tuple[list[tuple[Any, ...]], set[str]]:
    grouped: OrderedDict[str, list[tuple[Any, ...]]] = OrderedDict()
    for row in workbook["고객목록"].iter_rows(min_row=2, max_col=5, values_only=True):
        customer_id = clean(row[0])
        if customer_id:
            grouped.setdefault(customer_id, []).append(row)

    customers: list[tuple[Any, ...]] = []
    for customer_id, entries in grouped.items():
        names = [clean(row[1]) for row in entries if clean(row[1])]
        name = names[-1]
        notes = list(dict.fromkeys(clean(row[2]) for row in entries if clean(row[2])))
        previous_names = [value for value in dict.fromkeys(names) if value != name]
        notes.extend(f"이전 고객명: {value}" for value in previous_names)
        created_values = [timestamp_text(row[3]) for row in entries if timestamp_text(row[3])]
        updated_values = [timestamp_text(row[4]) for row in entries if timestamp_text(row[4])]
        created_at = min(created_values) if created_values else "2000-01-01 00:00:00"
        updated_at = max(updated_values + created_values) if updated_values or created_values else created_at
        customers.append((customer_id, name, "\n".join(notes), 0, created_at, updated_at))
    return customers, set(grouped)


def read_work_rows(workbook) -> list[tuple[str, int, tuple[Any, ...]]]:
    rows: list[tuple[str, int, tuple[Any, ...]]] = []
    for sheet_name, first_row in (("업무일지", 22), ("업무일지_과거", 3)):
        sheet = workbook[sheet_name]
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=first_row, max_col=16, values_only=True),
            first_row,
        ):
            if row[1] is not None:
                rows.append((sheet_name, row_number, row))
    return rows


def build_work_data(work_rows, known_customer_ids: set[str]):
    by_legacy: OrderedDict[str, list[tuple[str, int, tuple[Any, ...]]]] = OrderedDict()
    for record in work_rows:
        by_legacy.setdefault(clean(record[2][1]), []).append(record)

    work_logs: list[tuple[Any, ...]] = []
    details: list[tuple[Any, ...]] = []
    events: list[tuple[Any, ...]] = []
    orphan_first_date: dict[str, str] = {}

    for legacy_text, records in by_legacy.items():
        partitions: OrderedDict[tuple[str, str, str, str], list[tuple[str, int, tuple[Any, ...]]]] = OrderedDict()
        for record in records:
            row = record[2]
            shared = (date_text(row[3]) or "1900-01-01", clean(row[4]), clean(row[5]), clean(row[15]))
            partitions.setdefault(shared, []).append(record)

        split = len(partitions) > 1
        for part_number, (shared, part_rows) in enumerate(partitions.items(), 1):
            work_date, customer_id, work_type, content = shared
            work_id = f"excel-work-{legacy_text}" + (f"-{part_number}" if split else "")
            legacy_id = int(float(legacy_text))
            created_at = f"{work_date} 00:00:00"
            work_logs.append((work_id, legacy_id, work_date, customer_id, work_type, content, 0, created_at, created_at))
            if customer_id not in known_customer_ids:
                current = orphan_first_date.get(customer_id)
                orphan_first_date[customer_id] = min(current, work_date) if current else work_date

            for detail_number, (_, source_row_number, row) in enumerate(part_rows, 1):
                detail_id = f"excel-detail-{legacy_text}-{part_number}-{detail_number}-{source_row_number}"
                sequence = int(row[2]) if isinstance(row[2], (int, float)) else detail_number
                property_values = tuple(clean(row[index]) for index in range(6, 15))
                details.append((detail_id, work_id, sequence, *property_values))

                property_type, building_name, building_dong, unit_number, size_type, sale_price, jeonse_price, monthly_rent, source = property_values
                if work_type in EVENT_WORK_TYPES and property_type and building_name and unit_number:
                    key = listing_key((property_type, building_name, building_dong, unit_number))
                    events.append((
                        f"excel-event-{detail_id}", key, work_id, detail_id, work_date, work_type,
                        property_type, building_name, building_dong, unit_number, size_type,
                        sale_price, jeonse_price, monthly_rent, source, content, 0, created_at,
                    ))

    return work_logs, details, events, orphan_first_date


def build_listings(workbook) -> list[tuple[Any, ...]]:
    grouped: OrderedDict[str, list[tuple[int, tuple[Any, ...]]]] = OrderedDict()
    for row_number, row in enumerate(workbook["매물목록"].iter_rows(min_row=2, max_col=12, values_only=True), 2):
        if not any(value is not None for value in row):
            continue
        key = listing_key(row[3:7])
        grouped.setdefault(key, []).append((row_number, row))

    listings: list[tuple[Any, ...]] = []
    for key, records in grouped.items():
        def winner_rank(record):
            row_number, row = record
            active = date_text(row[1]) is None
            newest = date_text(row[0]) or date_text(row[1]) or ""
            return active, newest, row_number

        winner_number, winner = max(records, key=winner_rank)
        notes: list[str] = []
        winner_notes = clean(winner[11])
        if winner_notes:
            notes.append(winner_notes)
        if clean(winner[0]) == "NO매물":
            notes.append("[기존 등록일자: NO매물]")

        for row_number, row in records:
            if row_number == winner_number:
                continue
            previous = (
                f"[이전 매물 기록] 등록 {clean(row[0]) or '-'} / 말소 {clean(row[1]) or '-'} / "
                f"상태 {clean(row[2]) or '-'} / 타입 {clean(row[7]) or '-'} / "
                f"매매 {clean(row[8]) or '-'} / 전세 {clean(row[9]) or '-'} / 월세 {clean(row[10]) or '-'}"
            )
            notes.append(previous)
            if clean(row[11]):
                notes.append(clean(row[11]))

        registered_at = date_text(winner[0])
        closed_at = date_text(winner[1])
        updated_date = closed_at or registered_at or "2000-01-01"
        listings.append((
            f"excel-listing-{winner_number}", key, registered_at, closed_at, clean(winner[2]),
            clean(winner[3]), clean(winner[4]), clean(winner[5]), clean(winner[6]), clean(winner[7]),
            clean(winner[8]), clean(winner[9]), clean(winner[10]), "\n".join(notes), 0,
            f"{updated_date} 00:00:00",
        ))
    return listings


def build_migration(workbook_path: Path) -> tuple[str, dict[str, int]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    missing = REQUIRED_SHEETS - set(workbook.sheetnames)
    if missing:
        raise SystemExit(f"필수 시트가 없습니다: {', '.join(sorted(missing))}")

    customers, customer_ids = build_customers(workbook)
    work_rows = read_work_rows(workbook)
    work_logs, details, events, orphan_dates = build_work_data(work_rows, customer_ids)
    for customer_id, first_date in orphan_dates.items():
        created_at = f"{first_date} 00:00:00"
        customers.append((customer_id, customer_id, "엑셀 업무이력에서 자동 보완", 0, created_at, created_at))
    listings = build_listings(workbook)

    statements = [
        "DELETE FROM listing_events;",
        "DELETE FROM work_log_properties;",
        "DELETE FROM listings;",
        "DELETE FROM work_logs;",
        "DELETE FROM customers;",
    ]
    statements += insert_chunks("customers", ["id", "name", "notes", "is_demo", "created_at", "updated_at"], customers)
    statements += insert_chunks("work_logs", ["id", "legacy_id", "work_date", "customer_id", "work_type", "content", "is_demo", "created_at", "updated_at"], work_logs)
    statements += insert_chunks("work_log_properties", ["id", "work_log_id", "sequence", "property_type", "building_name", "building_dong", "unit_number", "size_type", "sale_price", "jeonse_price", "monthly_rent", "source"], details)
    statements += insert_chunks("listing_events", ["id", "listing_key", "work_log_id", "detail_id", "event_date", "status", "property_type", "building_name", "building_dong", "unit_number", "size_type", "sale_price", "jeonse_price", "monthly_rent", "source", "notes", "is_demo", "created_at"], events, 25)
    statements += insert_chunks("listings", ["id", "identity_key", "registered_at", "closed_at", "status", "property_type", "building_name", "building_dong", "unit_number", "size_type", "sale_price", "jeonse_price", "monthly_rent", "notes", "is_demo", "updated_at"], listings, 25)
    migration = "-- Private one-time Excel import. Do not commit this file to GitHub.\n" + BREAK.join(statements) + "\n"
    counts = {
        "customers": len(customers),
        "work_logs": len(work_logs),
        "work_log_properties": len(details),
        "listing_events": len(events),
        "listings": len(listings),
        "active_listings": sum(row[3] is None for row in listings),
        "orphan_customers_recovered": len(orphan_dates),
    }
    return migration, counts


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path, help="Excel에서 저장한 개발필요 .xlsx 파일")
    parser.add_argument("output", type=Path, help="생성할 비공개 SQL 마이그레이션")
    args = parser.parse_args()

    migration, counts = build_migration(args.workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(migration, encoding="utf-8")
    digest = hashlib.sha256(migration.encode("utf-8")).hexdigest()
    print("private migration built")
    for key, value in counts.items():
        print(f"{key}={value}")
    print(f"sha256={digest}")


if __name__ == "__main__":
    main()
